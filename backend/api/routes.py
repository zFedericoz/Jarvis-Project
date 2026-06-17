import asyncio
import base64
import json
import logging
import os
import platform
import subprocess
import re
import uuid
import time
from datetime import datetime, timezone
from collections import deque, defaultdict
from fastapi import APIRouter, WebSocket, UploadFile, File, Request, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from pathlib import Path

import psutil

from .dependencies import get_config, get_brain, get_speech, get_actions, get_memory, new_context, get_chat_manager
from .websocket_manager import manager, handle_wake_word, handle_audio_stream
import skills

logger = logging.getLogger("jarvis.api.routes")
UPLOAD_DIR = Path("/app/data/uploads")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
FEEDBACK_FILE = Path("/app/data/feedback.jsonl")

# ── Security Constants ────────────────────────────────────────────────────────
MAX_UPLOAD_SIZE = 50 * 1024 * 1024  # 50 MB
BLOCKED_EXTENSIONS = {'.env', '.key', '.pem', '.secret', '.db', '.git', '.cfg', '.ini', '.sql', '.pwd', '.pass'}
ALLOWED_EXTENSIONS = {'.txt', '.pdf', '.md', '.json', '.csv', '.log', '.py', '.js', '.ts', '.jsx', '.tsx',
                      '.yaml', '.yml', '.rst', '.html', '.css', '.xml', '.toml', '.docx', '.xlsx'}

# ── Rate Limiting Constants ───────────────────────────────────────────────────
MAX_CHAT_REQUESTS_PER_MINUTE = 10
_chat_request_times: defaultdict[str, deque] = defaultdict(lambda: deque(maxlen=100))

def _check_rate_limit(client_ip: str, max_per_minute: int = MAX_CHAT_REQUESTS_PER_MINUTE) -> bool:
    """Check if client exceeded rate limit (simple IP-based)"""
    now = time.time()
    times = _chat_request_times[client_ip]
    times.append(now)

    # Remove requests older than 1 minute
    while times and (now - times[0]) > 60:
        times.popleft()

    return len(times) <= max_per_minute

# ── Pending Actions Store (per azioni pericolose) ─────────────────────────────
_pending_actions: dict[str, dict] = {}
_PENDING_ACTION_TTL = 300  # 5 minutes

_DANGEROUS_ACTIONS = {"shutdown", "restart", "lock", "mute", "sleep", "hibernate"}

def _cleanup_expired_actions():
    """Remove pending actions that have timed out"""
    now = datetime.now(timezone.utc)
    expired = []
    for action_id, action_data in _pending_actions.items():
        created_at = datetime.fromisoformat(action_data.get("created_at", now.isoformat()))
        if (now - created_at).total_seconds() > _PENDING_ACTION_TTL:
            expired.append(action_id)
    for action_id in expired:
        del _pending_actions[action_id]
        logger.info(f"Pending action {action_id} expired")

def _is_dangerous(cmd: str) -> str | None:
    cmd_lower = cmd.lower()
    for kw in _DANGEROUS_ACTIONS:
        if kw in cmd_lower:
            return kw
    return None

# ── System Metrics Store ──────────────────────────────────────────────────────
_metric_history = {
    "cpu": deque(maxlen=60),
    "ram": deque(maxlen=60),
    "temp": deque(maxlen=60),
    "disk": deque(maxlen=60),
}

def _get_temperature():
    try:
        temps = psutil.sensors_temperatures()
        if temps:
            for name, entries in temps.items():
                if entries:
                    return round(entries[0].current, 1)
        if platform.system() == "Windows":
            result = subprocess.run(
                ["wmic", "/namespace:\\\\root\\wmi", "path", "MSAcpi_ThermalZoneTemperature", "get", "CurrentTemperature"],
                capture_output=True, text=True, timeout=5
            )
            match = re.search(r"(\d+)", result.stdout)
            if match:
                kelvin = int(match.group(1))
                return round(kelvin / 10 - 273.15, 1)
    except Exception:
        pass
    return None

def _get_top_processes(limit=5):
    procs = []
    for p in sorted(
        psutil.process_iter(["pid", "name", "cpu_percent", "memory_percent"]),
        key=lambda p: p.info.get("cpu_percent", 0) or 0,
        reverse=True,
    )[:limit]:
        try:
            procs.append({
                "pid": p.info["pid"],
                "name": p.info["name"],
                "cpu": round(p.info["cpu_percent"] or 0, 1),
                "mem": round(p.info["memory_percent"] or 0, 1),
            })
        except (psutil.NoSuchProcess, psutil.AccessDenied):
            continue
    return procs

def _collect_metrics():
    # Use non-blocking cpu_percent (interval=None)
    # First call requires interval for cache initialization
    cpu = psutil.cpu_percent(interval=None)
    ram = psutil.virtual_memory()
    disk = psutil.disk_usage("/")
    temp = _get_temperature()
    procs = _get_top_processes()
    net = psutil.net_io_counters()
    return {
        "cpu": round(cpu, 1),
        "ram": round(ram.percent, 1),
        "ram_gb": round(ram.used / (1024**3), 1),
        "ram_total_gb": round(ram.total / (1024**3), 1),
        "temp": temp,
        "disk": round(disk.percent, 1),
        "disk_gb": round(disk.used / (1024**3), 1),
        "disk_total_gb": round(disk.total / (1024**3), 1),
        "net_sent": round(net.bytes_sent / (1024**2), 2),
        "net_recv": round(net.bytes_recv / (1024**2), 2),
        "uptime": round(psutil.boot_time()),
        "processes": procs,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }

# ── System Event Logger ───────────────────────────────────────────────────────
_system_logs = deque(maxlen=100)
_system_logs.append({
    "timestamp": datetime.now(timezone.utc).isoformat(),
    "level": "info",
    "message": f"Sistema avviato — {platform.system()} {platform.release()}",
})

def _push_log(level: str, message: str):
    _system_logs.append({
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "level": level,
        "message": message,
    })

router = APIRouter(prefix="/api")

# ── Models ────────────────────────────────────────────────────────────────────

class ChatRequest(BaseModel):
    text: str = ""
    file_content: str = ""
    session_id: int | None = None
    stream: bool = False

class ChatResponse(BaseModel):
    response: str
    intent: str
    language: str
    audio: str | None = None
    session_id: int | None = None
    sources: list[dict] = []
    pending_action: str | None = None

class PendingActionInfo(BaseModel):
    id: str
    label: str
    description: str
    dangerous: bool = True

class ConfirmRequest(BaseModel):
    action_id: str
    confirm: bool

async def _sse_events(r: ChatResponse):
    d = {'type': 'done', 'response': r.response, 'intent': r.intent, 'language': r.language, 'session_id': r.session_id, 'sources': r.sources}
    if r.pending_action:
        d['pending_action'] = r.pending_action
    yield f"data: {json.dumps(d)}\n\n"

def _abort_response(session_id: int | None, stream: bool):
    r = ChatResponse(response="Richiesta interrotta.", intent="none", language="it", audio=None, session_id=session_id)
    if stream:
        return StreamingResponse(_sse_events(r), media_type="text/event-stream")
    return r

class StatusResponse(BaseModel):
    status: str
    version: str
    name: str
    llm: str

class UploadResponse(BaseModel):
    filename: str
    size: int
    content: str

class FeedbackRequest(BaseModel):
    message_id: str = ""
    user_message: str
    assistant_response: str
    rating: int
    language: str = "it"
    intent: str = "chat"

class GitRequest(BaseModel):
    command: str                       # es. "committa tutto", "status", "log"
    repo_path: str | None = None       # percorso repo; None = usa default da settings.yaml


# ── Standard routes ───────────────────────────────────────────────────────────

@router.get("/status", response_model=StatusResponse)
async def get_status():
    config = get_config()
    return StatusResponse(
        status="online",
        version="2.1.0",
        name="J.A.R.V.I.S.",
        llm=config["llm"]["model"],
    )

@router.get("/system/metrics")
async def get_system_metrics():
    metrics = _collect_metrics()
    _metric_history["cpu"].append(metrics["cpu"])
    _metric_history["ram"].append(metrics["ram"])
    if metrics["temp"] is not None:
        _metric_history["temp"].append(metrics["temp"])
    _metric_history["disk"].append(metrics["disk"])
    return {
        **metrics,
        "history": {
            "cpu": list(_metric_history["cpu"]),
            "ram": list(_metric_history["ram"]),
            "temp": list(_metric_history["temp"]),
            "disk": list(_metric_history["disk"]),
        }
    }

@router.get("/system/logs")
async def get_system_logs():
    cpu = psutil.cpu_percent(interval=0)
    ram = psutil.virtual_memory()
    if cpu > 85:
        _push_log("warning", f"CPU criticamente alta: {cpu}%")
    if ram.percent > 85:
        _push_log("warning", f"RAM criticamente alta: {ram.percent}%")
    if cpu < 10:
        _push_log("info", f"Sistema in idle — CPU {cpu}%, RAM {ram.percent}%")
    return {"logs": list(_system_logs)}

@router.websocket("/ws/wake")
async def websocket_wake(ws: WebSocket):
    await manager.connect(ws)
    config = get_config()
    await handle_wake_word(ws, config)
    manager.disconnect(ws)

@router.websocket("/ws/audio")
async def websocket_audio(ws: WebSocket):
    await manager.connect(ws)
    config = get_config()
    speech = get_speech(config)
    llm, intent_router, multiagent = get_brain(config)
    context = new_context()
    actions = get_actions(config)
    _, persistent_memory = get_memory(config)
    await handle_audio_stream(
        ws, speech["stt"], multiagent, intent_router,
        context, actions, speech["tts"], persistent_memory
    )
    manager.disconnect(ws)

@router.post("/chat")
async def chat_text(payload: ChatRequest, request: Request = None):
    # ── Rate limiting
    client_ip = request.client.host if request else "unknown"
    if not _check_rate_limit(client_ip):
        logger.warning(f"Rate limit exceeded for {client_ip}")
        raise HTTPException(status_code=429, detail=f"Rate limit: max {MAX_CHAT_REQUESTS_PER_MINUTE} requests/minute")

    config = get_config()
    text = payload.text.strip()
    file_content = payload.file_content
    session_id = payload.session_id

    if not text and not file_content:
        r = ChatResponse(response="No input provided.", intent="none", language="it", audio=None, session_id=session_id)
        if payload.stream:
            return StreamingResponse(_sse_events(r), media_type="text/event-stream")
        return r

    if file_content:
        text = f"{text}\n\n[File content]:\n{file_content}" if text else f"[File content]:\n{file_content}"

    # ── Chat session handling ────────────────────────────────────────────────
    chat_mgr = get_chat_manager()
    if session_id is not None:
        existing = chat_mgr.get_session(session_id)
        if not existing:
            session_id = None

    if session_id is None:
        session = chat_mgr.create_session()
        session_id = session["id"]

    chat_mgr.add_message(session_id, "user", payload.text.strip() or "[file]", "")

    llm, intent_router, multiagent = get_brain(config)
    context = new_context()
    actions = get_actions(config)
    speech = get_speech(config)
    _, persistent_memory = get_memory(config)

    if request and await request.is_disconnected():
        return _abort_response(session_id, payload.stream)

    lang = llm.detect_language(text)
    context.set_language(lang)
    original_query = text

    memories = persistent_memory.search(text, n_results=3)
    if memories:
        memory_context = "\n".join(f"Related memory: {m['text']}" for m in memories)
        prompt = f"{text}\n\n{memory_context}"
    else:
        prompt = text

    if request and await request.is_disconnected():
        return _abort_response(session_id, payload.stream)

    intent = intent_router.route(original_query)

    if intent in ("git", "terminal", "rpa", "productivity", "system_control", "media_player", "web_search"):
        if intent == "git":
            repo_path = persistent_memory.get_preference("cartella_progetti")
            response = await actions["git"].execute(original_query, repo_path=repo_path)
            _push_log("info", f"Git: {original_query[:60]}")
        elif intent == "web_search":
            try:
                loop = asyncio.get_event_loop()
                result = await loop.run_in_executor(None, lambda: skills.execute("web_search", query=original_query))
                response = f"🔍 Risultati ricerca per '{original_query}':\n{result}"
            except Exception:
                web_result, _ = multiagent._search_web(original_query)
                response = f"🔍 Risultati ricerca:\n{web_result}" if web_result else "Nessun risultato trovato."
        else:
            dangerous_key = _is_dangerous(original_query) if intent == "system_control" else None
            if dangerous_key:
                action_id = str(uuid.uuid4())
                _pending_actions[action_id] = {
                    "command": original_query,
                    "intent": intent,
                    "label": dangerous_key,
                    "created_at": datetime.now(timezone.utc).isoformat(),
                }
                r = ChatResponse(
                    response=f"⚠️ Richiesta azione pericolosa: `{dangerous_key}`. Attendo conferma.",
                    intent=intent, language=lang, audio=None,
                    session_id=session_id, pending_action=action_id,
                )
                if payload.stream:
                    return StreamingResponse(_sse_events(r), media_type="text/event-stream")
                return r

            response = await actions[intent].execute(original_query)

        if request and await request.is_disconnected():
            return _abort_response(session_id, payload.stream)

        context.add_turn("user", original_query)
        context.add_turn("assistant", response)
        persistent_memory.store(original_query, metadata={"role": "user", "intent": intent})
        persistent_memory.store(response, metadata={"role": "assistant", "intent": intent})
        chat_mgr.add_message(session_id, "assistant", response, intent)
        chat_mgr.auto_title(session_id)

        r = ChatResponse(response=response, intent=intent, language=lang, audio=None, session_id=session_id)
        if payload.stream:
            return StreamingResponse(_sse_events(r), media_type="text/event-stream")
        return r

    # ── Streaming LLM response ───────────────────────────────────────────────
    if payload.stream:
        async def stream_events():
            full_response = ""
            for token in multiagent.chat_stream(prompt, context.get_context(), language=lang, intent=intent, search_query=original_query):
                full_response += token
                yield f"data: {json.dumps({'type': 'token', 'text': token})}\n\n"
            sources = multiagent.last_sources
            done = {'type': 'done', 'response': full_response, 'intent': intent, 'language': lang, 'session_id': session_id, 'sources': sources}
            yield f"data: {json.dumps(done)}\n\n"

            context.add_turn("user", original_query)
            context.add_turn("assistant", full_response)
            persistent_memory.store(original_query, metadata={"role": "user", "intent": intent})
            persistent_memory.store(full_response, metadata={"role": "assistant", "intent": intent})
            chat_mgr.add_message(session_id, "assistant", full_response, intent)
            chat_mgr.auto_title(session_id)

        return StreamingResponse(stream_events(), media_type="text/event-stream")

    # ── Non-streaming LLM response ───────────────────────────────────────────
    response = multiagent.chat(
        prompt, context.get_context(),
        language=lang, intent=intent,
        search_query=original_query,
    )

    if request and await request.is_disconnected():
        return _abort_response(session_id, payload.stream)

    context.add_turn("user", original_query)
    context.add_turn("assistant", response)
    persistent_memory.store(original_query, metadata={"role": "user", "intent": intent})
    persistent_memory.store(response, metadata={"role": "assistant", "intent": intent})
    chat_mgr.add_message(session_id, "assistant", response, intent)
    chat_mgr.auto_title(session_id)

    return ChatResponse(response=response, intent=intent, language=lang, audio=None, session_id=session_id, sources=multiagent.last_sources)


@router.post("/confirm")
async def confirm_action(payload: ConfirmRequest):
    _cleanup_expired_actions()  # Clean up before checking

    action = _pending_actions.get(payload.action_id)
    if not action:
        return {"status": "error", "message": "Action not found or expired"}

    if payload.confirm:
        config = get_config()
        actions = get_actions(config)
        intent = action["intent"]
        if intent in actions:
            result = await actions[intent].execute(action["command"])
            del _pending_actions[payload.action_id]
            logger.info(f"Action confirmed: {payload.action_id}")
            return {"status": "ok", "result": result, "action": payload.action_id}
        return {"status": "error", "message": f"Action handler '{intent}' not found"}
    else:
        cmd = _pending_actions.pop(payload.action_id, None)
        logger.info(f"Action cancelled: {payload.action_id}")
        return {"status": "cancelled", "message": f"Action cancelled: {cmd['command'] if cmd else '?'}"}

@router.post("/upload", response_model=UploadResponse)
async def upload_file(file: UploadFile = File(...)):
    content = await file.read()

    # ── 1. Size validation
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=413, detail=f"File too large (max {MAX_UPLOAD_SIZE / (1024*1024):.0f} MB)")

    # ── 2. Sanitize filename (prevent path traversal)
    original_filename = file.filename or "upload"
    safe_filename = os.path.basename(original_filename)

    if not safe_filename or safe_filename in {'.', '..'}:
        safe_filename = str(uuid.uuid4())

    # ── 3. Extension validation
    file_ext = Path(safe_filename).suffix.lower()

    if file_ext in BLOCKED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type blocked for security: {file_ext}"
        )

    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type not allowed: {file_ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    # ── 4. Generate UUID-based filename to prevent collisions & path attacks
    safe_name = f"{uuid.uuid4()}{file_ext}"
    file_path = UPLOAD_DIR / safe_name
    file_path.parent.mkdir(parents=True, exist_ok=True)

    with open(file_path, "wb") as f:
        f.write(content)

    logger.info(f"File uploaded: {safe_name} (original: {original_filename}, size: {len(content)} bytes)")

    # ── 5. Extract text content
    text_content = ""
    try:
        if file_ext == ".docx":
            from docx import Document
            doc = Document(file_path)
            text_content = "\n".join(p.text for p in doc.paragraphs)
        elif file_ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            rows = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    rows.append("\t".join(str(c) if c is not None else "" for c in row))
            text_content = "\n".join(rows)
        elif file_ext in {".py", ".js", ".ts", ".jsx", ".tsx", ".yaml", ".yml", ".json", ".md", ".txt", ".rst", ".html", ".css", ".csv", ".xml", ".toml"}:
            text_content = content.decode("utf-8", errors="replace")
        else:
            text_content = content.decode("utf-8", errors="replace")
    except Exception as e:
        logger.warning(f"Text extraction failed for {safe_name}: {e}")
        text_content = f"[Unable to extract text: {e}]"

    return UploadResponse(filename=safe_name, size=len(content), content=text_content)


@router.post("/feedback")
async def submit_feedback(fb: FeedbackRequest):
    entry = {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "message_id": fb.message_id,
        "user_message": fb.user_message,
        "assistant_response": fb.assistant_response,
        "rating": fb.rating,
        "language": fb.language,
        "intent": fb.intent,
    }
    with open(FEEDBACK_FILE, "a", encoding="utf-8") as f:
        f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    logger.info(f"Feedback: rating={fb.rating}, lang={fb.language}, intent={fb.intent}")
    return {"status": "saved", "rating": fb.rating}


# ── Step 4: Git API routes ────────────────────────────────────────────────────

@router.post("/git/command")
async def git_command(payload: GitRequest):
    """
    Esegue un comando Git via API REST.

    Esempi:
      POST /api/git/command
      { "command": "committa tutto", "repo_path": "C:/Dev/JARVIS" }

      POST /api/git/command
      { "command": "status" }

      POST /api/git/command
      { "command": "ultimi 10 commit" }
    """
    config = get_config()
    actions = get_actions(config)
    git = actions.get("git")
    if not git:
        return {"status": "error", "message": "GitAction non disponibile"}

    # Se il repo_path non è specificato, prova dalle preferenze utente
    repo = payload.repo_path
    if not repo:
        _, persistent_memory = get_memory(config)
        repo = persistent_memory.get_preference("cartella_progetti")

    result = await git.execute(payload.command, repo_path=repo)
    _push_log("info", f"Git API: {payload.command[:60]}")
    return {"status": "ok", "result": result, "repo": str(repo or "default")}


@router.get("/git/status")
async def git_status(repo_path: str | None = None):
    """Shortcut per git status."""
    config = get_config()
    actions = get_actions(config)
    git = actions.get("git")
    if not git:
        return {"status": "error", "message": "GitAction non disponibile"}

    if not repo_path:
        _, mem = get_memory(config)
        repo_path = mem.get_preference("cartella_progetti")

    result = await git.execute("status", repo_path=repo_path)
    return {"status": "ok", "result": result}


@router.get("/git/log")
async def git_log(repo_path: str | None = None, n: int = 10):
    """Shortcut per git log con N commit."""
    config = get_config()
    actions = get_actions(config)
    git = actions.get("git")
    if not git:
        return {"status": "error", "message": "GitAction non disponibile"}

    if not repo_path:
        _, mem = get_memory(config)
        repo_path = mem.get_preference("cartella_progetti")

    result = await git.execute(f"mostrami gli ultimi {n} commit", repo_path=repo_path)
    return {"status": "ok", "result": result, "n": n}


@router.post("/git/commit")
async def git_commit(repo_path: str | None = None):
    """
    Esegue git add -A + commit con messaggio generato dall'LLM.
    Shortcut per il frontend (pulsante "Commit" nella dashboard).
    """
    config = get_config()
    actions = get_actions(config)
    git = actions.get("git")
    if not git:
        return {"status": "error", "message": "GitAction non disponibile"}

    if not repo_path:
        _, mem = get_memory(config)
        repo_path = mem.get_preference("cartella_progetti")

    result = await git.execute("committa tutto", repo_path=repo_path)
    return {"status": "ok", "result": result}


@router.post("/git/push")
async def git_push(repo_path: str | None = None):
    """Esegue git push."""
    config = get_config()
    actions = get_actions(config)
    git = actions.get("git")
    if not git:
        return {"status": "error", "message": "GitAction non disponibile"}

    if not repo_path:
        _, mem = get_memory(config)
        repo_path = mem.get_preference("cartella_progetti")

    result = await git.execute("push", repo_path=repo_path)
    return {"status": "ok", "result": result}

# ══════════════════════════════════════════════════════════════════════════════
# Step 5 — Terminal API routes
# ══════════════════════════════════════════════════════════════════════════════

class TerminalRequest(BaseModel):
    command: str          # comando shell oppure frase in linguaggio naturale
    cwd: str | None = None  # working directory; None = usa terminal.working_dir da config


@router.post("/terminal/run")
async def terminal_run(payload: TerminalRequest):
    """
    Esegue un comando nel terminale sicuro.

    Esempi:
      { "command": "ls -la" }
      { "command": "che versione di Python ho?" }
      { "command": "esegui python analisi.py", "cwd": "/app/scripts" }
      { "command": "docker ps" }
    """
    config = get_config()
    actions = get_actions(config)
    terminal = actions.get("terminal")
    if not terminal:
        return {"status": "error", "message": "TerminalAction non disponibile"}

    result = await terminal.execute(payload.command, cwd=payload.cwd)
    _push_log("info", f"Terminal: {payload.command[:60]}")
    return {
        "status": "ok",
        "command": payload.command,
        "result": result,
    }


@router.get("/terminal/history")
async def terminal_history():
    """Ritorna lo storico dei comandi eseguiti (ultimi 50, più recenti prima)."""
    config = get_config()
    actions = get_actions(config)
    terminal = actions.get("terminal")
    if not terminal:
        return {"status": "error", "message": "TerminalAction non disponibile"}

    return {
        "status": "ok",
        "history": terminal.get_history(),
    }


@router.get("/terminal/allowed")
async def terminal_allowed():
    """
    Ritorna la mappa dei comandi consentiti nelle categorie abilitate.
    Utile per il frontend (mostrare all'utente cosa può fare).
    """
    config = get_config()
    actions = get_actions(config)
    terminal = actions.get("terminal")
    if not terminal:
        return {"status": "error", "message": "TerminalAction non disponibile"}

    return {
        "status": "ok",
        "allowed": terminal.get_allowed_commands(),
        "working_dir": str(terminal._working_dir),
        "timeout": terminal._timeout,
    }

# ══════════════════════════════════════════════════════════════════════════════
# Step 6 — Focus Mode API routes
# Aggiungere in fondo a backend/api/routes.py
# ══════════════════════════════════════════════════════════════════════════════

class FocusStartRequest(BaseModel):
    duration_minutes: int | None = None   # None = usa il default da settings.yaml (25 min)

class FocusSiteRequest(BaseModel):
    site: str                             # es. "reddit.com"


@router.post("/focus/start")
async def focus_start(payload: FocusStartRequest = FocusStartRequest()):
    """
    Avvia una sessione Pomodoro.
    Blocca i siti distraenti e attiva Focus Assist (Windows).

    Esempio:
      POST /api/focus/start
      POST /api/focus/start  { "duration_minutes": 45 }
    """
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    result = productivity.focus.start(work_minutes=payload.duration_minutes)
    _push_log("info", f"Focus avviato: {payload.duration_minutes or 'default'} min")
    return {"status": "ok", "result": result, "focus_status": productivity.focus.status()}


@router.post("/focus/stop")
async def focus_stop():
    """Ferma la sessione focus e sblocca i siti."""
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    result = productivity.focus.stop()
    _push_log("info", "Focus disattivato")
    return {"status": "ok", "result": result}


@router.post("/focus/pause")
async def focus_pause():
    """Mette in pausa il timer focus (sblocca i siti temporaneamente)."""
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    result = productivity.focus.pause()
    return {"status": "ok", "result": result}


@router.post("/focus/resume")
async def focus_resume():
    """Riprende una sessione focus in pausa."""
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    result = productivity.focus.resume()
    return {"status": "ok", "result": result}


@router.get("/focus/status")
async def focus_status():
    """
    Ritorna lo stato corrente della sessione focus.

    Response:
      {
        "state": "working" | "break" | "paused" | "idle",
        "pomodoro_count": 3,
        "remaining_seconds": 847,
        "remaining_formatted": "14 min 7 sec",
        "blocked_sites": ["reddit.com", ...],
        "notifications_muted": true
      }
    """
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    return {"status": "ok", **productivity.focus.status()}


@router.post("/focus/sites/add")
async def focus_add_site(payload: FocusSiteRequest):
    """
    Aggiunge un sito alla blacklist focus (persistente tra sessioni).

    Esempio:
      POST /api/focus/sites/add  { "site": "reddit.com" }
    """
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    result = productivity.focus.add_site(payload.site)
    return {"status": "ok", "result": result, "blocked_sites": productivity.focus._blocked_sites}


@router.post("/focus/sites/remove")
async def focus_remove_site(payload: FocusSiteRequest):
    """
    Rimuove un sito dalla blacklist focus.

    Esempio:
      POST /api/focus/sites/remove  { "site": "youtube.com" }
    """
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    result = productivity.focus.remove_site(payload.site)
    return {"status": "ok", "result": result, "blocked_sites": productivity.focus._blocked_sites}


@router.get("/focus/sites")
async def focus_list_sites():
    """Ritorna la lista dei siti bloccati durante il focus."""
    config = get_config()
    actions = get_actions(config)
    productivity = actions.get("productivity")
    if not productivity:
        return {"status": "error", "message": "Productivity action non disponibile"}

    return {
        "status": "ok",
        "blocked_sites": productivity.focus._blocked_sites,
        "count": len(productivity.focus._blocked_sites),
    }

# ══════════════════════════════════════════════════════════════════════════════
# Step 7 — RPA / UI Automation API routes
# ══════════════════════════════════════════════════════════════════════════════

class RPAClickRequest(BaseModel):
    x: int
    y: int
    button: str = "left"       # "left" | "right" | "middle"
    clicks: int = 1            # 1 = singolo, 2 = doppio

class RPATypeRequest(BaseModel):
    text: str
    interval: float = 0.03    # pausa tra tasti in secondi

class RPAHotkeyRequest(BaseModel):
    keys: list[str]            # es. ["ctrl", "s"] oppure ["alt", "tab"]

class RPAOpenAppRequest(BaseModel):
    app: str                   # nome app (es. "chrome", "vscode", "notepad")

class RPAOpenFileRequest(BaseModel):
    path: str                  # percorso file
    app: str | None = None     # app da usare (None = predefinita del sistema)

class RPAScrollRequest(BaseModel):
    direction: str = "down"    # "up" | "down"
    clicks: int = 3

class RPADragRequest(BaseModel):
    x1: int
    y1: int
    x2: int
    y2: int

class RPACommandRequest(BaseModel):
    command: str               # comando in linguaggio naturale


@router.post("/rpa/screenshot")
async def rpa_screenshot():
    """
    Cattura uno screenshot del desktop e lo salva in data/screenshots/.

    Response:
      { "status": "ok", "result": "Screenshot salvato: data/screenshots/screenshot_20260101_120000.png" }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, rpa._do_screenshot)
    _push_log("info", "Screenshot catturato")
    return {"status": "ok", "result": result}


@router.post("/rpa/analyze")
async def rpa_analyze():
    """
    Cattura lo schermo e lo analizza con il LLM vision.
    Richiede un modello con supporto immagini (llava, qwen2-vl, ecc.).
    Se il modello non supporta vision, ritorna una descrizione base.
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    result = await rpa._do_analyze("analizza lo schermo")
    _push_log("info", "Analisi schermo completata")
    return {"status": "ok", "result": result}


@router.post("/rpa/click")
async def rpa_click(payload: RPAClickRequest):
    """
    Esegue un click del mouse alle coordinate specificate.

    Esempio:
      POST /api/rpa/click  { "x": 500, "y": 300, "button": "left", "clicks": 1 }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    gui = rpa._gui()
    if not gui:
        return {"status": "error", "message": "pyautogui non disponibile"}

    loop = asyncio.get_event_loop()
    await loop.run_in_executor(
        None,
        lambda: gui.click(payload.x, payload.y, button=payload.button,
                          clicks=payload.clicks, duration=rpa._move_duration)
    )
    _push_log("info", f"Click {payload.button} su ({payload.x},{payload.y})")
    return {"status": "ok", "result": f"Click su ({payload.x},{payload.y}) eseguito"}


@router.post("/rpa/type")
async def rpa_type(payload: RPATypeRequest):
    """
    Digita testo nella finestra attiva.

    Esempio:
      POST /api/rpa/type  { "text": "ciao mondo", "interval": 0.05 }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    gui = rpa._gui()
    if not gui:
        return {"status": "error", "message": "pyautogui non disponibile"}

    loop = asyncio.get_event_loop()
    await loop.run_in_executor(None, lambda: gui.write(payload.text, interval=payload.interval))
    _push_log("info", f"Digitato: {payload.text[:40]}")
    return {"status": "ok", "result": f"Digitato: \"{payload.text}\""}


@router.post("/rpa/hotkey")
async def rpa_hotkey(payload: RPAHotkeyRequest):
    """
    Preme una combinazione di tasti.

    Esempi:
      { "keys": ["ctrl", "s"] }        → Salva
      { "keys": ["alt", "tab"] }       → Cambia finestra
      { "keys": ["ctrl", "shift", "t"]} → Riapri tab chiusa
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    gui = rpa._gui()
    if not gui:
        return {"status": "error", "message": "pyautogui non disponibile"}

    loop = asyncio.get_event_loop()
    await loop.run_in_executor(None, lambda: gui.hotkey(*payload.keys))
    combo = "+".join(k.upper() for k in payload.keys)
    _push_log("info", f"Hotkey: {combo}")
    return {"status": "ok", "result": f"Premuto: {combo}"}


@router.post("/rpa/open_app")
async def rpa_open_app(payload: RPAOpenAppRequest):
    """
    Apre un'applicazione per nome.

    Esempi:
      { "app": "chrome" }
      { "app": "vscode" }
      { "app": "notepad" }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, rpa._do_open_app, f"apri {payload.app}")
    _push_log("info", f"App aperta: {payload.app}")
    return {"status": "ok", "result": result}


@router.post("/rpa/open_file")
async def rpa_open_file(payload: RPAOpenFileRequest):
    """
    Apre un file con l'app specificata o quella predefinita.

    Esempi:
      { "path": "C:/Dev/main.py", "app": "code" }
      { "path": "/home/user/report.pdf" }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    app_hint = f" con {payload.app}" if payload.app else ""
    cmd = f"apri il file {payload.path}{app_hint}"
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, rpa._do_open_file, cmd)
    return {"status": "ok", "result": result}


@router.post("/rpa/scroll")
async def rpa_scroll(payload: RPAScrollRequest):
    """
    Esegue lo scroll nella finestra attiva.

    Esempio:
      { "direction": "down", "clicks": 5 }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    cmd = f"scorri in {payload.direction} di {payload.clicks} click"
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, rpa._do_scroll, cmd)
    return {"status": "ok", "result": result}


@router.post("/rpa/drag")
async def rpa_drag(payload: RPADragRequest):
    """
    Esegue un drag & drop tra due coordinate.

    Esempio:
      { "x1": 100, "y1": 200, "x2": 400, "y2": 300 }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    cmd = f"trascina da ({payload.x1},{payload.y1}) a ({payload.x2},{payload.y2})"
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, rpa._do_drag, cmd)
    return {"status": "ok", "result": result}


@router.get("/rpa/screen_info")
async def rpa_screen_info():
    """Ritorna risoluzione schermo e posizione corrente del mouse."""
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    result = rpa._do_screen_info()
    return {"status": "ok", "result": result}


@router.post("/rpa/command")
async def rpa_command(payload: RPACommandRequest):
    """
    Esegue un comando RPA in linguaggio naturale.
    Equivalente a parlare direttamente a JARVIS.

    Esempi:
      { "command": "apri Chrome" }
      { "command": "premi Ctrl+S" }
      { "command": "cosa c'è sullo schermo?" }
      { "command": "clicca su (500, 300)" }
    """
    config = get_config()
    actions = get_actions(config)
    rpa = actions.get("rpa")
    if not rpa:
        return {"status": "error", "message": "RPAAction non disponibile"}

    result = await rpa.execute(payload.command)
    _push_log("info", f"RPA command: {payload.command[:60]}")
    return {"status": "ok", "result": result}


# ══════════════════════════════════════════════════════════════════════════════
# Step 8 — Chat Session routes
# ══════════════════════════════════════════════════════════════════════════════

class RenameSessionRequest(BaseModel):
    title: str


@router.get("/chats")
async def list_chat_sessions():
    chat_mgr = get_chat_manager()
    return {"sessions": chat_mgr.list_sessions()}


@router.post("/chats")
async def create_chat_session():
    chat_mgr = get_chat_manager()
    session = chat_mgr.create_session()
    return {"session": session}


@router.delete("/chats/{session_id}")
async def delete_chat_session(session_id: int):
    chat_mgr = get_chat_manager()
    ok = chat_mgr.delete_session(session_id)
    return {"status": "deleted" if ok else "not_found"}


@router.patch("/chats/{session_id}")
async def rename_chat_session(session_id: int, payload: RenameSessionRequest):
    chat_mgr = get_chat_manager()
    ok = chat_mgr.rename_session(session_id, payload.title)
    return {"status": "renamed" if ok else "not_found"}


@router.get("/chats/{session_id}/messages")
async def get_chat_messages(session_id: int):
    chat_mgr = get_chat_manager()
    messages = chat_mgr.get_messages(session_id)
    return {"messages": messages}


# ══════════════════════════════════════════════════════════════════════════════
# Market Data API routes
# ══════════════════════════════════════════════════════════════════════════════

@router.get("/market/quote")
async def market_quote(symbol: str = "AAPL"):
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, lambda: skills.execute("market_data", action="quote", symbol=symbol))
    try:
        return json.loads(result)
    except (json.JSONDecodeError, TypeError):
        return {"error": result}


@router.get("/market/history")
async def market_history(symbol: str = "AAPL", period: str = "1mo", interval: str = "1d"):
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, lambda: skills.execute("market_data", action="history", symbol=symbol, period=period, interval=interval))
    try:
        return json.loads(result)
    except (json.JSONDecodeError, TypeError):
        return {"error": result, "data": []}


@router.get("/market/search")
async def market_search(q: str = "Tesla"):
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, lambda: skills.execute("market_data", action="search", query=q))
    try:
        return json.loads(result)
    except (json.JSONDecodeError, TypeError):
        return {"error": result, "results": []}


@router.get("/market/news")
async def market_news(symbol: str = "AAPL"):
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, lambda: skills.execute("market_data", action="news", symbol=symbol))
    try:
        return json.loads(result)
    except (json.JSONDecodeError, TypeError):
        return {"error": result}
