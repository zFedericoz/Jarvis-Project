import io
import logging
import os
import uuid
from pathlib import Path

import psutil
from fastapi import APIRouter, File, HTTPException, UploadFile, WebSocket
from fastapi.responses import Response, StreamingResponse

from .constants import MAX_UPLOAD_SIZE, BLOCKED_EXTENSIONS, ALLOWED_EXTENSIONS
from .dependencies import get_config, get_brain, get_speech, get_actions, get_memory, new_context, get_chat_manager
from .routes_common import (
    _collect_metrics, _metric_history, _system_logs, _push_log,
    UPLOAD_DIR,
)
from .schemas import StatusResponse, UploadResponse, RagThresholdRequest, ExportRequest, VoiceSettingsRequest
from plugins.loader import get_plugin_loader
from .websocket_manager import manager, handle_wake_word, handle_audio_stream
from brain.rag_searcher import rag_distance_threshold

logger = logging.getLogger("jarvis.api.routes")

router = APIRouter()


@router.get("/health")
async def health():
    """Basic alive check — sempre 200 se il server è in esecuzione."""
    return {"status": "ok"}


@router.get("/ready")
async def ready():
    """Verifica che le dipendenze siano pronte."""
    deps = {"server": True}
    try:
        get_chat_manager()
        deps["database"] = True
    except Exception as e:
        deps["database"] = str(e)
    try:
        import os
        import redis.asyncio as aioredis
        r = aioredis.Redis(host="redis", port=6379, password=os.getenv("REDIS_PASSWORD"), decode_responses=True, socket_timeout=2)
        await r.ping()
        deps["redis"] = True
    except Exception:
        deps["redis"] = "unreachable"
    try:
        from chromadb import HttpClient
        client = HttpClient(host="chromadb", port=8000, settings={"allow_reset": True})
        client.heartbeat()
        deps["chromadb"] = True
    except Exception:
        deps["chromadb"] = "unreachable"
    try:
        llm, _, _ = get_brain(get_config())
        deps["llm"] = True
    except Exception as e:
        deps["llm"] = str(e)

    all_ok = all(v is True for v in deps.values())
    return {"status": "ok" if all_ok else "degraded", "checks": deps}


@router.get("/status", response_model=StatusResponse)
async def get_status():
    config = get_config()
    return StatusResponse(
        status="online",
        version="2.1.0",
        name="J.A.R.V.I.S.",
        llm=config["llm"]["model"],
    )


@router.get("/system/metrics")
async def get_system_metrics():
    metrics = _collect_metrics()
    _metric_history["cpu"].append(metrics["cpu"])
    _metric_history["ram"].append(metrics["ram"])
    if metrics["temp"] is not None:
        _metric_history["temp"].append(metrics["temp"])
    _metric_history["disk"].append(metrics["disk"])
    return {
        **metrics,
        "history": {
            "cpu": list(_metric_history["cpu"]),
            "ram": list(_metric_history["ram"]),
            "temp": list(_metric_history["temp"]),
            "disk": list(_metric_history["disk"]),
        }
    }


@router.get("/system/logs")
async def get_system_logs():
    cpu = psutil.cpu_percent(interval=0)
    ram = psutil.virtual_memory()
    if cpu > 85:
        _push_log("warning", f"CPU criticamente alta: {cpu}%")
    if ram.percent > 85:
        _push_log("warning", f"RAM criticamente alta: {ram.percent}%")
    if cpu < 10:
        _push_log("info", f"Sistema in idle — CPU {cpu}%, RAM {ram.percent}%")
    return {"logs": list(_system_logs)}


@router.websocket("/ws/wake")
async def websocket_wake(ws: WebSocket):
    await manager.connect(ws)
    config = get_config()
    await handle_wake_word(ws, config)
    manager.disconnect(ws)


@router.websocket("/ws/audio")
async def websocket_audio(ws: WebSocket):
    await manager.connect(ws)
    config = get_config()
    speech = get_speech(config)
    llm, intent_router, multiagent = get_brain(config)
    context = new_context()
    actions = get_actions(config)
    _, persistent_memory = get_memory(config)
    await handle_audio_stream(
        ws, speech["stt"], multiagent, intent_router,
        context, actions, speech["tts"], persistent_memory
    )
    manager.disconnect(ws)


@router.get("/rag/threshold")
async def get_rag_threshold():
    return {"threshold": rag_distance_threshold}


@router.put("/rag/threshold")
async def set_rag_threshold(body: RagThresholdRequest):
    global rag_distance_threshold
    if not 0.1 <= body.threshold <= 3.0:
        raise HTTPException(status_code=400, detail="Threshold must be between 0.1 and 3.0")
    rag_distance_threshold = body.threshold
    logger.info(f"RAG distance threshold changed to {body.threshold}")
    return {"threshold": rag_distance_threshold}


def _build_pdf(messages: list[dict]) -> bytes:
    from fpdf import FPDF
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_font("DejaVu", "", r"/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", uni=True)
    pdf.add_font("DejaVu", "B", r"/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", uni=True)
    pdf.set_font("DejaVu", "B", 16)
    pdf.cell(0, 10, "J.A.R.V.I.S. - Conversazione", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(5)
    for m in messages:
        role = "TU" if m["role"] == "user" else "J.A.R.V.I.S."
        c = (0, 0, 0) if m["role"] == "user" else (0, 128, 0)
        pdf.set_font("DejaVu", "B", 10)
        pdf.set_text_color(*c)
        pdf.cell(0, 6, role, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 9)
        pdf.set_text_color(30, 30, 30)
        text = m.get("text", "")
        text = text.encode("latin-1", errors="replace").decode("latin-1")
        pdf.multi_cell(0, 5, text)
        pdf.ln(2)
    return bytes(pdf.output())


@router.post("/export/pdf")
async def export_pdf(body: ExportRequest):
    try:
        pdf_bytes = _build_pdf([{"role": m.role, "text": m.text} for m in body.messages])
        return Response(content=pdf_bytes, media_type="application/pdf",
                        headers={"Content-Disposition": "attachment; filename=jarvis-export.pdf"})
    except Exception as e:
        logger.error(f"PDF export fallito: {e}")
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {e}")


def _load_voice_settings():
    import yaml
    cfg_path = Path("config/settings.yaml")
    if not cfg_path.exists():
        return {"wake_word_enabled": True, "wake_word_sensitivity": 0.5, "stt_model": "base", "stt_language": "auto", "tts_engine": "kokoro", "tts_speed": 1.0}
    with open(cfg_path) as f:
        cfg = yaml.safe_load(f)
    s = cfg.get("speech", {})
    ww = s.get("wake_word", {})
    return {
        "wake_word_enabled": ww.get("enabled", True),
        "wake_word_sensitivity": ww.get("sensitivity", 0.5),
        "stt_model": s.get("stt", {}).get("model", "base"),
        "stt_language": s.get("stt", {}).get("language", "auto"),
        "tts_engine": s.get("tts", {}).get("engine", "kokoro"),
        "tts_speed": s.get("tts", {}).get("kokoro_speed", 1.0),
    }


@router.get("/voice/settings")
async def get_voice_settings():
    return _load_voice_settings()


@router.put("/voice/settings")
async def update_voice_settings(body: VoiceSettingsRequest):
    import yaml
    cfg_path = Path("config/settings.yaml")
    cfg = {}
    if cfg_path.exists():
        with open(cfg_path) as f:
            cfg = yaml.safe_load(f) or {}
    if "speech" not in cfg:
        cfg["speech"] = {}
    s = cfg["speech"]
    if body.wake_word_enabled is not None:
        s.setdefault("wake_word", {})["enabled"] = body.wake_word_enabled
    if body.wake_word_sensitivity is not None:
        s.setdefault("wake_word", {})["sensitivity"] = body.wake_word_sensitivity
    if body.stt_model is not None:
        s.setdefault("stt", {})["model"] = body.stt_model
    if body.stt_language is not None:
        s.setdefault("stt", {})["language"] = body.stt_language
    if body.tts_engine is not None:
        s.setdefault("tts", {})["engine"] = body.tts_engine
    if body.tts_speed is not None:
        s.setdefault("tts", {})["kokoro_speed"] = body.tts_speed
    cfg_path.parent.mkdir(parents=True, exist_ok=True)
    with open(cfg_path, "w") as f:
        yaml.dump(cfg, f, default_flow_style=False)
    logger.info(f"Voice settings aggiornati")
    return _load_voice_settings()


from pydantic import BaseModel

class AuthRequest(BaseModel):
    username: str
    password: str

@router.post("/auth/register")
async def register(body: AuthRequest):
    from .auth import register_user
    return register_user(body.username, body.password)

@router.post("/auth/login")
async def login(body: AuthRequest):
    from .auth import authenticate_user
    return authenticate_user(body.username, body.password)


@router.get("/plugins")
async def list_plugins():
    loader = get_plugin_loader()
    return {"plugins": loader.list_plugins()}


@router.post("/plugins/{name}/exec")
async def exec_plugin(name: str, body: dict):
    loader = get_plugin_loader()
    plugin = loader.get_plugin(name)
    if not plugin:
        raise HTTPException(status_code=404, detail=f"Plugin '{name}' not found")
    action = body.get("action", "execute")
    params = body.get("params", {})
    result = await plugin.execute(action, params)
    return {"plugin": name, "result": result, "action": action}


@router.get("/vectors")
async def get_vectors():
    """Restituisce tutti i vettori e i testi delle collezioni ChromaDB per visualizzazione."""
    try:
        from memory.persistent import PersistentMemory
        pm = PersistentMemory()
        all_data = {}
        for col_name, col in [("memories", pm.collection), ("knowledge", pm.knowledge), ("preferences", pm.prefs)]:
            data = col.get(include=["documents", "metadatas", "embeddings"])
            ids = data.get("ids")
            if ids is not None and len(ids) > 0:
                pts = []
                embeddings = data.get("embeddings")
                docs = data.get("documents") or []
                metas = data.get("metadatas") or []
                for i in range(len(ids)):
                    emb = None
                    if isinstance(embeddings, (list, tuple)) and i < len(embeddings) and embeddings[i] is not None:
                        emb = list(embeddings[i])[:3]
                    elif hasattr(embeddings, "__getitem__") and embeddings[i] is not None:
                        emb = list(embeddings[i])[:3]
                    pts.append({
                        "id": ids[i],
                        "text": docs[i] if i < len(docs) else "",
                        "metadata": metas[i] if i < len(metas) else {},
                        "embedding": emb
                    })
                all_data[col_name] = pts
        return {"collections": all_data, "total": sum(len(v) for v in all_data.values())}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/upload", response_model=UploadResponse)
async def upload_file(file: UploadFile = File(...)):
    content = await file.read()

    # ── 1. Size validation
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=413, detail=f"File too large (max {MAX_UPLOAD_SIZE / (1024*1024):.0f} MB)")

    # ── 2. Sanitize filename (prevent path traversal)
    original_filename = file.filename or "upload"
    safe_filename = os.path.basename(original_filename)

    if not safe_filename or safe_filename in {'.', '..'}:
        safe_filename = str(uuid.uuid4())

    # ── 3. Extension validation
    file_ext = Path(safe_filename).suffix.lower()

    if file_ext in BLOCKED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type blocked for security: {file_ext}"
        )

    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type not allowed: {file_ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    # ── 4. Generate UUID-based filename to prevent collisions & path attacks
    safe_name = f"{uuid.uuid4()}{file_ext}"
    file_path = UPLOAD_DIR / safe_name
    file_path.parent.mkdir(parents=True, exist_ok=True)

    with open(file_path, "wb") as f:
        f.write(content)

    logger.info(f"File uploaded: {safe_name} (original: {original_filename}, size: {len(content)} bytes)")

    # ── 5. Extract text content
    text_content = ""
    try:
        if file_ext == ".docx":
            from docx import Document
            doc = Document(file_path)
            text_content = "\n".join(p.text for p in doc.paragraphs)
        elif file_ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            rows = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    rows.append("\t".join(str(c) if c is not None else "" for c in row))
            text_content = "\n".join(rows)
        elif file_ext in {".py", ".js", ".ts", ".jsx", ".tsx", ".yaml", ".yml", ".json", ".md", ".txt", ".rst", ".html", ".css", ".csv", ".xml", ".toml"}:
            text_content = content.decode("utf-8", errors="replace")
        else:
            text_content = content.decode("utf-8", errors="replace")
    except Exception as e:
        logger.warning(f"Text extraction failed for {safe_name}: {e}")
        text_content = f"[Unable to extract text: {e}]"

    return UploadResponse(filename=safe_name, size=len(content), content=text_content)
